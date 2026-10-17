"""Rehearse the whole closing workflow end to end, against a COPY of the database.

    set MOP_RECO_DB=D:\\MOP_Reco\\db\\rehearsal.sqlite
    python scripts/round_trip_rehearsal.py

It refuses to run against db/reco.sqlite. Using the SAME functions the Exceptions page calls:

  1. download the Exception bills file (the table's own export), type ONE transaction number in the Txn
     column, upload the same file, preview it, apply it;
  2. trace the effect: the bill and its transaction leave the exception tables, both appear in Manual match
     bills / transactions, the store task closes, the Summary receivable falls by the amount, the controls hold;
  3. reverse it and show everything returned (and the task raised again);
  4. download the Exception transactions file, fill in a customer refund on one row, upload, trace the same
     things - and show it never closed a bill and never moved a receivable;
  5. reverse the refund.

Every step prints what it found. Files are written to exports/round_trip/ so they can be opened.
"""

from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "app"))

DB = os.environ.get("MOP_RECO_DB", "")
if not DB or Path(DB).name == "reco.sqlite":
    sys.exit("Set MOP_RECO_DB to a COPY of the database (never db/reco.sqlite).")

from mop_reco import exception_pairing as ep  # noqa: E402
from mop_reco import exception_tables as et  # noqa: E402
from mop_reco import manual_match as mm  # noqa: E402
from mop_reco import queries as q  # noqa: E402
from mop_reco import refunds as rf  # noqa: E402
from mop_reco import store_sheet_source as src  # noqa: E402
from mop_reco import table_upload as tu  # noqa: E402
from mop_reco.db import (Exception_, RefundEvent, Run, StoreTask, TransactionRefund, get_engine,  # noqa: E402
                         get_session, init_db)
from mop_reco.export import ExportMeta, build_export  # noqa: E402

RUN, A, B = 11, date(2026, 4, 1), date(2026, 8, 31)
WHO = "Round-trip rehearsal"
OUT = ROOT / "exports" / "round_trip"
OUT.mkdir(parents=True, exist_ok=True)
CONFIG = yaml.safe_load((ROOT / "app" / "config.yaml").read_text(encoding="utf-8"))
MINUTES, DAYS = ep.window_settings(CONFIG)
BILL_TYPES = q.NOT_RECEIVED_TYPES
TXN_TYPES = q.RECEIVED_NOT_BILLED_TYPES + ("original_txn_not_found", "store_ambiguous")

engine = init_db(get_engine())
session = get_session(engine)
print(f"database: {DB}")


def say(text=""):
    print(text, flush=True)


def figures(label, store=None, bill_no=None):
    ov = q.match_overview(session, RUN, A, B)
    mat = src.mop_particulars_matrix(session, RUN, A, B)
    sheet, _ = src.load_fast(session, RUN, A, B, store) if store else (None, None)
    task = None
    if bill_no:
        t = (session.query(StoreTask).filter(StoreTask.task_type == "open_bill", StoreTask.bill_no == bill_no)
             .order_by(StoreTask.generation.desc()).first())
        task = None if t is None else (t.generation, t.status, t.closed_by, t.close_reason)
    f = {
        "label": label,
        "exception bills (open)": (ov.exception_bills_n, round(ov.exception_bills_v, 2)),
        "exception transactions (open)": (ov.exception_txns_n, round(ov.exception_txns_v, 2)),
        "bills open/auto/manual/total": tuple(ov.bills[k]["count"] for k in ("open", "auto", "manual", "total")),
        "txns auto/manual/refunded/unmapped/total": (ov.txns.auto_n, ov.txns.manual_n, ov.txns.refunded_n,
                                                      ov.txns.unmapped_n, ov.txns.total_n),
        "controls": ("bills OK" if ov.bills_ok else "bills MISMATCH",
                     "txns OK" if ov.txns_ok else "txns MISMATCH",
                     "bills=exceptions OK" if ov.bills_agree_with_exceptions else "bills=exceptions MISMATCH",
                     "unmapped=exceptions OK" if ov.unmapped_agrees_with_exceptions else "unmapped=exceptions MISMATCH"),
        "Summary total-sale receivable (Rs)": round(mat.total_sale.receivable, 2),
        "Summary digital receivable, that store (Rs)": None if sheet is None else round(float(sheet["digital_receivable"].sum()), 2),
        "store task": task,
    }
    say(f"--- {label}")
    for k, v in f.items():
        if k != "label":
            say(f"    {k}: {v}")
    return f


# ---------------------------------------------------------------------------------------------
say("=" * 100)
say("PART 1 - MATCH A BILL TO A TRANSACTION BY UPLOADING THE TABLE'S OWN FILE")
say("=" * 100)
pool = ep.unmapped_pool(session)
pop = ep.open_population(session, RUN)
assignment = ep.assign(pop, pool, MINUTES, DAYS, ep.mop_family(CONFIG))
live_tasks = {b for (b,) in session.query(StoreTask.bill_no).filter(StoreTask.task_type == "open_bill",
                                                                     StoreTask.status.in_(("open", "replied", "under_review"))).all()}
pick = next(r for r in pop.sort_values(["amount", "bill_no"], ascending=[False, True]).itertuples()
            if assignment.kind.get(r.bill_key) == "exact" and 800 <= r.amount <= 3000 and r.bill_no in live_tasks)
cand = assignment.candidates[pick.bill_key][0]
say(f"bill chosen: {pick.bill_no}  store {pick.store}  {pick.date}  {pick.mop}  Rs {pick.amount:,.2f}")
say(f"the system's own suggestion for it: transaction {cand.txn_id} ({cand.gateway}) Rs {cand.amount:,.2f}")
before = figures("BEFORE", pick.store, pick.bill_no)

say("\n[1] download Exception bills - the table's export")
shape = lambda f: et.add_pairing(et.bill_frame(session, RUN, f), assignment)[0]  # noqa: E731
frame = q.exception_export(session, A, B, None, BILL_TYPES, ("open",), transform=shape)
cols = et.bill_columns(pairing=True, screen=False)
res = build_export(et.export_view(frame, cols, "Exception bills"),
                   ExportMeta(run_id=RUN, run_timestamp=session.get(Run, RUN).started_at,
                              date_from=A, date_to=B, filters={"show": "Open"}, value_column="amount",
                              title="Exception bills"))
downloaded = OUT / "Exception_bills_downloaded.xlsx"
downloaded.write_bytes(res.data)
say(f"    {len(frame):,} rows, columns exactly: {cols}")
untouched = tu.preview_bills(session, RUN, tu.read_table_file(res.data, res.filename, cols))
say(f"    (for the record: uploading that file UNTOUCHED would match {len(untouched.accepted):,} bills - "
    f"the system's own suggestions - and reject {len(untouched.rejected):,}; {len(untouched.ignored):,} have nothing filled)")

say("\n[2] someone fills in ONE transaction number in the Txn 1 column and saves the file")
wb = openpyxl.load_workbook(downloaded)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows())
head = next(i for i, r in enumerate(rows) if r[0].value == "date")
col = {str(c.value): c.column for c in rows[head] if c.value}
for r in rows[head + 1:]:
    for name in ("Txn 1", "Txn 2", "Amount 1", "Amount 2"):
        if name in col:
            r[col[name] - 1].value = None
    if str(r[col["bill_no"] - 1].value) == pick.bill_no:
        r[col["Txn 1"] - 1].value = str(cand.txn_id)
        r[col["Txn 1"] - 1].number_format = "@"
filled = OUT / "Exception_bills_filled.xlsx"
wb.save(filled)
say(f"    typed {cand.txn_id} against bill {pick.bill_no}; file saved: {filled}")

say("\n[3] upload the same file - preview")
up_df = tu.read_table_file(filled.read_bytes(), filled.name, cols)
prev = tu.preview_bills(session, RUN, up_df)
say(f"    {len(prev.accepted)} accepted, {len(prev.rejected)} rejected, {len(prev.ignored):,} with no change")
for r in prev.accepted + prev.rejected:
    say(f"    row {r.row_no}: {r.kind} {r.ref} -> {r.counterpart}: {r.verdict} {r.why}")

say("\n[4] confirm - apply (all or nothing)")
counts = tu.apply_preview(session, prev, WHO, file_name=filled.name)
say(f"    applied: {counts}")
key_bill = mm.bill_key_of({"business_date": pick.date, "store": pick.store, "bill_no": pick.bill_no, "mop": pick.mop})
after = figures("AFTER THE MATCH", pick.store, pick.bill_no)
mb = et.manual_bill_frame(session, RUN)
mt = et.manual_txn_frame(session)
say(f"    bill exception status: {session.query(Exception_.status).filter_by(exception_key=key_bill).scalar()}")
say(f"    transaction exception status: {session.query(Exception_.status).filter_by(exception_key=f'{cand.gateway}|{cand.txn_id}').scalar()}")
say(f"    in Exception bills (open)? {bool(len(q.exception_page(session, A, B, None, BILL_TYPES, ('open',), limit=1, offset=0, sort='Amount (high to low)')[0].query('exception_key == @key_bill')))}")
say("    Manual match bills row:")
say(mb[mb["bill_no"] == pick.bill_no][["date", "store", "bill_no", "amount", "Txn 1", "matched_txn", "matched_by", "matched_at", "source"]].to_string(index=False))
say("    Manual match transactions row:")
say(mt[mt["transaction_id"] == str(cand.txn_id)][["gateway", "transaction_id", "store", "amount", "matched_bill", "resolution", "matched_by", "matched_at", "source"]].to_string(index=False))
say(f"    Summary receivable fell by: Rs {before['Summary total-sale receivable (Rs)'] - after['Summary total-sale receivable (Rs)']:,.2f} "
    f"(that store: Rs {before['Summary digital receivable, that store (Rs)'] - after['Summary digital receivable, that store (Rs)']:,.2f}; the bill was Rs {pick.amount:,.2f})")
say(f"    store task: {before['store task']}  ->  {after['store task']}")

say("\n[5] reverse the match")
entry = session.query(mm.ManualMatch).filter_by(bill_key=key_bill, lifecycle="active").one()
mm.reverse(session, entry.id, WHO, "round-trip rehearsal - reversing")
rev = figures("AFTER REVERSAL", pick.store, pick.bill_no)
say(f"    everything back as before (except the new task generation)? "
    f"{all(before[k] == rev[k] for k in before if k not in ('label', 'store task'))}")
say(f"    store task: {before['store task']}  ->  {rev['store task']}")
say(f"    the reversed match is kept: lifecycle={session.get(mm.ManualMatch, entry.id).lifecycle}, by {session.get(mm.ManualMatch, entry.id).reversed_by}")

# ---------------------------------------------------------------------------------------------
say()
say("=" * 100)
say("PART 2 - A CUSTOMER REFUND RECORDED BY UPLOADING THE TABLE'S OWN FILE")
say("=" * 100)
suggested = {c.txn_id for cl in assignment.candidates.values() for c in cl}
pool2 = pool[~pool["txn_id"].isin(suggested)].sort_values(["amount", "txn_id"], ascending=[False, True])
# a transaction whose reference NO bill holds (80 duplicate Axis lines share a reference a bill already took -
# those can be neither matched nor refunded; Part 3 shows that refusal)
t = next(r for r in pool2.itertuples() if 500 <= r.amount <= 3000 and rf._owner_of(session, r.gateway, str(r.txn_id)) is None)
say(f"transaction chosen (no bill is suggested for it): {t.txn_id} ({t.gateway}) store {t.store} Rs {t.amount:,.2f}")
before2 = figures("BEFORE", t.store)

say("\n[1] download Exception transactions - the table's export")
tframe = q.exception_export(session, A, B, None, TXN_TYPES, ("open",), transform=lambda f: et.txn_frame(session, f))
tcols = list(et.TXN_COLUMNS)
tres = build_export(et.export_view(tframe, tcols, "Exception transactions"),
                    ExportMeta(run_id=RUN, run_timestamp=None, date_from=A, date_to=B, filters={"show": "Open"},
                               value_column="amount", title="Exception transactions"))
tdown = OUT / "Exception_transactions_downloaded.xlsx"
tdown.write_bytes(tres.data)
say(f"    {len(tframe):,} rows, columns exactly: {tcols}")

say("\n[2] fill customer_name, zoho_no and refund_date on ONE row and save")
wb = openpyxl.load_workbook(tdown)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows())
head = next(i for i, r in enumerate(rows) if r[0].value == "gateway")
col = {str(c.value): c.column for c in rows[head] if c.value}
for r in rows[head + 1:]:
    if str(r[col["transaction_id"] - 1].value) == str(t.txn_id):
        r[col["customer_name"] - 1].value = "Ramesh Kumar"
        r[col["zoho_no"] - 1].value = "ZOHO-REHEARSAL-1"
        r[col["refund_date"] - 1].value = "05/09/2026"
tfilled = OUT / "Exception_transactions_filled.xlsx"
wb.save(tfilled)
say(f"    filled the refund fields against transaction {t.txn_id}; file saved: {tfilled}")

say("\n[3] upload the same file - preview")
tprev = tu.preview_txns(session, RUN, tu.read_table_file(tfilled.read_bytes(), tfilled.name, tcols))
say(f"    {len(tprev.accepted)} accepted, {len(tprev.rejected)} rejected, {len(tprev.ignored):,} with no change")
for r in tprev.accepted + tprev.rejected:
    say(f"    row {r.row_no}: {r.kind} {r.ref} -> {r.counterpart}: {r.verdict} {r.why}")

say("\n[4] confirm - apply")
say(f"    applied: {tu.apply_preview(session, tprev, WHO, file_name=tfilled.name)}")
after2 = figures("AFTER THE REFUND", t.store)
tkey = f"{t.gateway}|{t.txn_id}"
say(f"    transaction exception status: {session.query(Exception_.status).filter_by(exception_key=tkey).scalar()}  "
    f"({session.query(Exception_.reason_code).filter_by(exception_key=tkey).scalar()})")
mt2 = et.manual_txn_frame(session)
say("    Manual match transactions row:")
say(mt2[mt2["transaction_id"] == str(t.txn_id)][["gateway", "transaction_id", "store", "amount", "customer_name", "zoho_no", "refund_date", "resolution", "matched_by", "matched_at", "source"]].to_string(index=False))
say(f"    bills untouched: open bills {before2['bills open/auto/manual/total']} -> {after2['bills open/auto/manual/total']}")
say(f"    revenue untouched: Summary receivable {before2['Summary total-sale receivable (Rs)']} -> {after2['Summary total-sale receivable (Rs)']}")

say("\n[5] reverse the refund")
rid = session.query(TransactionRefund).filter_by(transaction_ref=str(t.txn_id)).one().id
rf.reverse_refund(session, rid, WHO, "round-trip rehearsal - reversing")
rev2 = figures("AFTER REVERSAL", t.store)
say(f"    everything back as before? {all(before2[k] == rev2[k] for k in before2 if k not in ('label', 'store task'))}")
say(f"    refund events kept: {[(e.action, e.by) for e in session.query(RefundEvent).filter_by(transaction_ref=str(t.txn_id)).order_by(RefundEvent.id).all()]}")


# ---------------------------------------------------------------------------------------------
say()
say("=" * 100)
say("PART 3 - THE REFUSALS, ON REAL DATA")
say("=" * 100)
dup = next(r for r in pool.sort_values("txn_id").itertuples() if rf._owner_of(session, r.gateway, str(r.txn_id)) is not None)
owner = rf._owner_of(session, dup.gateway, str(dup.txn_id))
say(f"transaction {dup.txn_id} ({dup.gateway}, store {dup.store}, Rs {dup.amount:,.2f}) is listed as an exception, but its "
    f"reference is already held by bill {owner.bill_no}")
other = next(r for r in pop.itertuples() if r.store == dup.store and r.bill_no != owner.bill_no)
say(f"(a) an upload matching open bill {other.bill_no} to it:")
base = tu.read_table_file(downloaded.read_bytes(), downloaded.name, cols).assign(**{"Txn 1": "", "Txn 2": ""})
one = base.copy()
one.loc[one["bill_no"] == other.bill_no, "Txn 1"] = str(dup.txn_id)
bf = tu.preview_bills(session, RUN, one)
for r in bf.rows:
    if r.verdict != "ignore":
        say(f"    row {r.row_no}: {r.kind} {r.ref} -> {r.counterpart}: {r.verdict.upper()} {r.why}")
say("(b) all or nothing - one good row and one bad row in the same file:")
mixed = base.copy()
mixed.loc[mixed["bill_no"] == pick.bill_no, "Txn 1"] = str(cand.txn_id)
mixed.loc[mixed["bill_no"] == other.bill_no, "Txn 1"] = str(dup.txn_id)
mprev = tu.preview_bills(session, RUN, mixed)
say(f"    preview: {len(mprev.accepted)} accepted, {len(mprev.rejected)} rejected -> can apply: {mprev.can_apply}")
n_before = session.query(mm.ManualMatch).filter_by(lifecycle="active").count()
try:
    tu.apply_preview(session, mprev, WHO)
except ValueError as exc:
    say(f"    apply refused: {exc}")
say(f"    active manual matches before {n_before}, after {session.query(mm.ManualMatch).filter_by(lifecycle='active').count()} - nothing written")
