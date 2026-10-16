"""CMS Reco - one block per DEPOSIT DATE, chained per store (2026-10-16).

Replaces the running-ledger layout (Opening receivable / Days outstanding). The operator's own
model, as drawn in the reference workbook. Per store, per deposit date `d`:

    OPENING          = the previous block's CLOSING  (the first block: a configured opening, default 0)
    Cash/CMS Deposit = -1 x SUM(PTC "PTC Bill Amount SUM") where store, Billdate = d, PTC Head = Cash/CMS Deposit
    CMS Report (1)   = SUM(Axis SLIPAMOUNT) where store, DEPDATE = d AND CRDATE = d        and TYPE = COL
    CMS Report (2)   = SUM(Axis SLIPAMOUNT) where store, DEPDATE = d AND CRDATE = d + 1    and TYPE = COL
    CLOSING          = OPENING + Deposit - CMS(1) - CMS(2)

Two decisions that came from the data, not the brief:

* The credit date is CRDATE. CLGDATE equals DEPDATE on every one of 23,902 COL rows (and on every
  ADJ row), so it cannot separate a same-day credit from a next-day one.
* A credit dated day+2 or later falls in NEITHER column. It is never dropped: it stays in CLOSING
  (the deposit genuinely had not been credited by day+1) and is listed separately by
  `credited_later()` so it can be chased.

The store is NOT in the Axis file (there is no "Store Code" column - column BP of the operator's
workbook is a pasted lookup). It comes from the pickup-code master, location-aware (APM/MMR share
AX31918). A row that does not resolve is store = UNRESOLVED and stays in every total; BNK is its own
pseudo-store, never folded into a real store.

Pure of Streamlit (CLAUDE.md section 8). Reads only the rows `persist()` wrote at run time.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

from . import masters
from .db import (CmsPtcDeposit, CmsRecoRemark, CmsRecoRemarkEvent, CmsSlip, StoreDayAgg)

UNRESOLVED = "UNRESOLVED"
PSEUDO_STORES = ("BNK", UNRESOLVED)          # shown after the real stores, never merged into one
EPS = 0.005
LATE_FLAG_DAYS = 10                           # a credit this many days (or more) after deposit is flagged by name

# Axis CMS columns by POSITION (a header lookup would silently keep one of the duplicated names).
_POS = {"TYPE": 0, "LOCAT": 3, "DEPDATE": 7, "CRDATE": 10, "SLIPNO": 15, "ADJ": 20, "SLIPAMOUNT": 22,
        "RECOVEREDAMOUNT": 23, "PICKUPPOINTCODE": 66}
_ORDER = sorted(_POS, key=_POS.get)


# ---------------------------------------------------------------------------
# Reading the raw Axis CMS files (all of them, including credits dated after the run's last date)
# ---------------------------------------------------------------------------

def _iso(value):
    if value is None or value is pd.NaT or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text or text.lower() in ("nan", "nat", "none"):
        return None
    for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    return None if pd.isna(parsed) else parsed.strftime("%Y-%m-%d")


def _read_file(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".zip":
        with zipfile.ZipFile(path) as zf:
            member = next(n for n in zf.namelist() if n.lower().endswith(".csv"))
            raw = zf.read(member)
        df = pd.read_csv(io.BytesIO(raw), header=0, dtype=str, index_col=False,
                         usecols=[_POS[c] for c in _ORDER], encoding="utf-8")
        df.columns = _ORDER                      # usecols returns them in file order
    elif suffix in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = wb[wb.sheetnames[0]].iter_rows(min_row=2, values_only=True)
            df = pd.DataFrame([[r[_POS[c]] if len(r) > _POS[c] else None for c in _ORDER] for r in rows], columns=_ORDER)
        finally:
            wb.close()
    else:
        raise ValueError(f"{path.name}: not a CMS file type this reader knows")
    df["source_file"] = path.name
    df["_mtime"] = path.stat().st_mtime
    return df


def read_cms_rows(config) -> pd.DataFrame:
    """Every Axis CMS row from every discovered file, one row per slip. A slip present in several
    extractions resolves to the most recently modified file (CLAUDE.md section 4)."""
    files = list(config.source_status("axis_cms").files)
    frames = [_read_file(Path(f)) for f in files if Path(f).suffix.lower() in (".zip", ".xlsx", ".xlsm")]
    if not frames:
        return pd.DataFrame(columns=_ORDER)
    df = pd.concat(frames, ignore_index=True)
    df["TYPE"] = df["TYPE"].astype(str).str.strip()
    df["dep_date"] = df["DEPDATE"].map(_iso)
    df["cr_date"] = df["CRDATE"].map(_iso)
    for src, dst in (("SLIPAMOUNT", "amount"), ("ADJ", "adj"), ("RECOVEREDAMOUNT", "recovered")):
        df[dst] = pd.to_numeric(df[src].astype(str).str.replace(",", ""), errors="coerce")
    df["pickup_code"] = df["PICKUPPOINTCODE"].astype(str).str.strip().replace({"nan": None, "": None})
    df["location"] = df["LOCAT"].astype(str).str.strip().replace({"nan": None, "": None})
    df["slip_no"] = df["SLIPNO"].astype(str).str.strip().replace({"nan": None, "": None})
    df = df.sort_values("_mtime")
    # natural key: TYPE + pickup + deposit date + slip no (an ADJ row names the slip it adjusts, so
    # TYPE is part of the key). ZPO carries no deposit date: slip, amount and credit date stand in.
    df["_key"] = (df["TYPE"] + "|" + df["pickup_code"].fillna("") + "|" + df["dep_date"].fillna("") + "|"
                  + df["slip_no"].fillna("") + "|" + np.where(df["dep_date"].isna(),
                                                              df["amount"].astype(str) + "|" + df["cr_date"].fillna(""), ""))
    df = df.drop_duplicates("_key", keep="last").reset_index(drop=True)
    return df[df["TYPE"].isin(["COL", "ADJ.CR.", "ADJ.DB.", "ZPO"])]


def resolve_stores(session, rows: pd.DataFrame) -> pd.DataFrame:
    """Store per row via the pickup-code master (location-aware). Never guessed: anything that does not
    resolve to exactly one store is UNRESOLVED."""
    rows = rows.copy()
    cache: dict = {}
    out = []
    for r in rows.itertuples(index=False):
        if r.TYPE == "ZPO" or not r.pickup_code:
            out.append(UNRESOLVED)
            continue
        key = (r.pickup_code, r.location, r.dep_date)
        if key not in cache:
            as_of = date.fromisoformat(r.dep_date) if r.dep_date else date.today()
            res = masters.resolve_pickup_store(session, r.pickup_code, r.location, as_of)
            cache[key] = res.store if (res.status == "resolved" and res.store) else UNRESOLVED
        out.append(cache[key])
    rows["store"] = out
    return rows


def read_ptc_deposits(warehouse: Path, date_from: str, date_to: str) -> pd.DataFrame:
    """PTC Cash/CMS Deposit per store per business date, carried positive (source x -1)."""
    folder = Path(warehouse) / "ptc"
    frames = []
    for p in sorted(folder.glob("2???-??-??.parquet")):
        if not (date_from <= p.stem <= date_to):
            continue
        d = pd.read_parquet(p, columns=["store", "business_date", "is_deposit", "amount"])
        d = d[d["is_deposit"].fillna(False).astype(bool)]
        if not d.empty:
            frames.append(d)
    if not frames:
        return pd.DataFrame(columns=["dep_date", "store", "amount"])
    d = pd.concat(frames, ignore_index=True)
    d["amount"] = -pd.to_numeric(d["amount"], errors="coerce").fillna(0.0)     # PAYMENT rows are negative
    d = d.rename(columns={"business_date": "dep_date"})
    return d.groupby(["dep_date", "store"], as_index=False)["amount"].sum()


# ---------------------------------------------------------------------------
# Persisting (a run stage) and loading (the page)
# ---------------------------------------------------------------------------

def run_range(session, run_id: int):
    lo = session.query(StoreDayAgg.business_date).filter(StoreDayAgg.run_id == run_id).order_by(StoreDayAgg.business_date).first()
    hi = (session.query(StoreDayAgg.business_date).filter(StoreDayAgg.run_id == run_id)
          .order_by(StoreDayAgg.business_date.desc()).first())
    return (lo[0], hi[0]) if lo else (None, None)


def persist(session, run_id: int, config, warehouse: Path = None) -> dict:
    """Write the CMS slips and PTC deposits this run's page reads. Replaces any earlier rows for the run."""
    warehouse = Path(warehouse) if warehouse else Path(__file__).resolve().parents[2] / "warehouse"
    lo, hi = run_range(session, run_id)
    if lo is None:
        return {"slips": 0, "deposits": 0}
    rows = resolve_stores(session, read_cms_rows(config))
    keep = rows[(rows["TYPE"] == "ZPO") | ((rows["dep_date"] >= lo) & (rows["dep_date"] <= hi))].copy()
    dep = read_ptc_deposits(warehouse, lo, hi)

    session.query(CmsSlip).filter(CmsSlip.run_id == run_id).delete()
    session.query(CmsPtcDeposit).filter(CmsPtcDeposit.run_id == run_id).delete()
    if not keep.empty:
        session.execute(CmsSlip.__table__.insert(), [
            {"run_id": run_id, "row_type": r.TYPE, "dep_date": r.dep_date, "cr_date": r.cr_date, "store": r.store,
             "pickup_code": r.pickup_code, "location": r.location, "slip_no": r.slip_no,
             "amount": 0.0 if pd.isna(r.amount) else float(r.amount),
             "adj": None if pd.isna(r.adj) else float(r.adj),
             "recovered": None if pd.isna(r.recovered) else float(r.recovered), "source_file": r.source_file}
            for r in keep.itertuples(index=False)])
    if not dep.empty:
        session.execute(CmsPtcDeposit.__table__.insert(), [
            {"run_id": run_id, "dep_date": r.dep_date, "store": r.store, "amount": float(r.amount)}
            for r in dep.itertuples(index=False)])
    session.commit()
    return {"slips": int(len(keep)), "deposits": int(len(dep)), "first": lo, "last": hi}


def load(session, run_id: int):
    """(slips, deposits) frames for the run."""
    slips = pd.read_sql(CmsSlip.__table__.select().where(CmsSlip.run_id == run_id), session.bind)
    deposits = pd.read_sql(CmsPtcDeposit.__table__.select().where(CmsPtcDeposit.run_id == run_id), session.bind)
    for c in ("amount", "adj", "recovered"):
        if c in slips:
            slips[c] = slips[c].astype(float)
    if "amount" in deposits:
        deposits["amount"] = deposits["amount"].astype(float)
    return slips, deposits


def opening_from_config(config_raw: dict) -> dict:
    """{store: opening CMS outstanding at the first loaded date}. Empty = zero for every store; a file
    can be wired in later by filling `cms_reco.opening` - no code change."""
    section = ((config_raw or {}).get("cms_reco") or {}).get("opening") or {}
    return {str(k): float(v) for k, v in section.items()}


# ---------------------------------------------------------------------------
# The model - blocks chained per store
# ---------------------------------------------------------------------------

@dataclass
class Model:
    dates: list                      # ISO strings, consecutive calendar days
    stores: list
    deposit: np.ndarray              # [date, store]
    cms1: np.ndarray
    cms2: np.ndarray
    opening: np.ndarray
    closing: np.ndarray
    first_opening: dict = field(default_factory=dict)

    def idx(self, d: str) -> int:
        return self.dates.index(d)

    def window(self, date_from: str, date_to: str):
        ids = [i for i, d in enumerate(self.dates) if date_from <= d <= date_to]
        return (ids[0], ids[-1] + 1) if ids else (0, 0)


def _order_stores(stores) -> list:
    real = sorted(s for s in set(stores) if s not in PSEUDO_STORES)
    return real + [s for s in PSEUDO_STORES if s in set(stores)]


def build_model(slips: pd.DataFrame, deposits: pd.DataFrame, first: str, last: str, openings: dict = None) -> Model:
    """The chained blocks over every calendar day from `first` to `last`, built block by block."""
    openings = openings or {}
    d0, d1 = date.fromisoformat(first), date.fromisoformat(last)
    dates = [(d0 + timedelta(days=i)).isoformat() for i in range((d1 - d0).days + 1)]
    col = slips[slips["row_type"] == "COL"] if not slips.empty else slips
    stores = _order_stores(list(deposits["store"]) + (list(col["store"]) if not col.empty else []) + list(openings))
    sx = {s: i for i, s in enumerate(stores)}
    dx = {d: i for i, d in enumerate(dates)}
    n_d, n_s = len(dates), len(stores)
    deposit, cms1, cms2 = (np.zeros((n_d, n_s)) for _ in range(3))
    for r in deposits.itertuples(index=False):
        if r.dep_date in dx and r.store in sx:
            deposit[dx[r.dep_date], sx[r.store]] += r.amount
    if not col.empty:
        cd = pd.to_datetime(col["dep_date"], errors="coerce")
        cr = pd.to_datetime(col["cr_date"], errors="coerce")
        lag = (cr - cd).dt.days
        for dep_date, store, amount, g in zip(col["dep_date"], col["store"], col["amount"], lag):
            if dep_date not in dx or store not in sx or pd.isna(g):
                continue
            if g == 0:
                cms1[dx[dep_date], sx[store]] += amount
            elif g == 1:
                cms2[dx[dep_date], sx[store]] += amount
    deposit, cms1, cms2 = (np.round(a, 2) for a in (deposit, cms1, cms2))
    opening, closing = np.zeros((n_d, n_s)), np.zeros((n_d, n_s))
    first_open = np.array([openings.get(s, 0.0) for s in stores], dtype=float)
    prev = first_open
    for i in range(n_d):
        opening[i] = prev
        closing[i] = np.round(opening[i] + deposit[i] - cms1[i] - cms2[i], 2)
        prev = closing[i]
    return Model(dates, stores, deposit, cms1, cms2, opening, closing, dict(zip(stores, first_open)))


def summary_frame(m: Model, date_from: str, date_to: str) -> pd.DataFrame:
    """STORE | Opening | Cash/CMS Deposit | CMS Report | Diff. Opening is the FIRST block's opening in the
    window (the chain's own, not zero); Diff = Opening + Deposit - CMS Report."""
    a, b = m.window(date_from, date_to)
    if a == b:
        return pd.DataFrame(columns=["store", "opening", "deposit", "cms", "diff"])
    dep = m.deposit[a:b].sum(axis=0)
    cms = (m.cms1[a:b] + m.cms2[a:b]).sum(axis=0)
    opening = m.opening[a]
    return pd.DataFrame({"store": m.stores, "opening": np.round(opening, 2), "deposit": np.round(dep, 2),
                         "cms": np.round(cms, 2), "diff": np.round(opening + dep - cms, 2)})


def blocks_frame(m: Model, date_from: str, date_to: str) -> pd.DataFrame:
    """The wide date-block detail for a window: STORE then, per date, Opening | Deposit | CMS (1) | CMS (2) | Closing."""
    a, b = m.window(date_from, date_to)
    cols = {"STORE NAME": m.stores}
    for i in range(a, b):
        tag = date.fromisoformat(m.dates[i]).strftime("%d/%m")
        cols[f"{tag} Opening"] = m.opening[i]
        cols[f"{tag} Deposit"] = m.deposit[i]
        cols[f"{tag} CMS (same day)"] = m.cms1[i]
        cols[f"{tag} CMS (next day)"] = m.cms2[i]
        cols[f"{tag} Closing"] = m.closing[i]
    return pd.DataFrame(cols)


# ---------------------------------------------------------------------------
# Controls
# ---------------------------------------------------------------------------

@dataclass
class Control:
    label: str
    ok: bool
    detail: str = ""
    amount: float = 0.0


def block_controls(m: Model, date_from: str, date_to: str) -> list:
    """The three structural controls, recomputed independently of how the model was built."""
    a, b = m.window(date_from, date_to)
    out = []
    if a == b:
        return out
    recomputed = np.round(m.opening[a:b] + m.deposit[a:b] - m.cms1[a:b] - m.cms2[a:b], 2)
    bad = np.abs(recomputed - m.closing[a:b]) > EPS
    out.append(Control("Every block: CLOSING = OPENING + Deposit - CMS(1) - CMS(2)", not bad.any(),
                       f"{int(bad.sum())} store-block(s) differ" if bad.any() else "",
                       float(np.abs(recomputed - m.closing[a:b])[bad].sum()) if bad.any() else 0.0))
    chain_bad = np.abs(m.opening[a + 1:b] - m.closing[a:b - 1]) > EPS if b - a > 1 else np.zeros((0, len(m.stores)), bool)
    out.append(Control("Every chain: block N+1 OPENING = block N CLOSING, per store", not chain_bad.any(),
                       f"{int(chain_bad.sum())} link(s) broken" if chain_bad.any() else "",
                       float(np.abs(m.opening[a + 1:b] - m.closing[a:b - 1])[chain_bad].sum()) if chain_bad.any() else 0.0))
    s = summary_frame(m, date_from, date_to)
    diff_bad = np.abs(s["diff"].to_numpy() - m.closing[b - 1]) > EPS
    out.append(Control("Summary Diff = the last block's CLOSING, per store", not diff_bad.any(),
                       f"{int(diff_bad.sum())} store(s) differ" if diff_bad.any() else "",
                       float(np.abs(s["diff"].to_numpy() - m.closing[b - 1])[diff_bad].sum()) if diff_bad.any() else 0.0))
    return out


def deposit_tie(session, run_id: int, deposits: pd.DataFrame, date_from: str, date_to: str) -> Control:
    """The page's PTC deposits against the store sheet's column I over the same dates (two separate paths)."""
    from sqlalchemy import func
    page = float(deposits[(deposits["dep_date"] >= date_from) & (deposits["dep_date"] <= date_to)]["amount"].sum())
    sheet = float(session.query(func.coalesce(func.sum(StoreDayAgg.cash_deposit), 0)).filter(
        StoreDayAgg.run_id == run_id, StoreDayAgg.business_date >= date_from,
        StoreDayAgg.business_date <= date_to).scalar() or 0)
    diff = round(page - sheet, 2)
    return Control("Page Cash/CMS Deposit ties to store sheet column I", abs(diff) <= EPS,
                   f"page {page:,.2f} against store sheet {sheet:,.2f}", diff)


@dataclass
class Bridge:
    page_cms: float
    credited_later: float
    all_col_by_deposit_date: float
    after_cutoff: float
    unmapped: float
    adj_net: float
    expected_j: float
    store_sheet_j: float
    unexplained: float
    cutoff: str

    @property
    def difference(self) -> float:
        return round(self.page_cms - self.store_sheet_j, 2)


def bridge_to_j(session, run_id: int, m: Model, slips: pd.DataFrame, date_from: str, date_to: str) -> Bridge:
    """Why the page's CMS Report total is not store-sheet column J - every rupee of it, named.

    J is built on the DEPOSIT date from COL (SLIPAMOUNT) + ADJ.CR. (ADJ) - ADJ.DB. (RECOVEREDAMOUNT), for
    stores that resolve, from credits the run had loaded - i.e. dated up to the run's last date. The page is
    COL only, credits dated day or day+1 only, every store (UNRESOLVED and BNK included). So:

        page CMS Report
      + credited later than day+1                  -> every COL slip by deposit date
      - credits dated after the run's last date    (not in J's source)
      - rows whose store does not resolve          (J drops them)
      + ADJ.CR. - ADJ.DB. net                      (J includes them, the page does not)
      = what J should be - compared with J itself; any rest is `unexplained`."""
    from sqlalchemy import func
    _lo, cutoff = run_range(session, run_id)
    a, b = m.window(date_from, date_to)
    page_cms = float((m.cms1[a:b] + m.cms2[a:b]).sum())
    w = slips[(slips["dep_date"] >= date_from) & (slips["dep_date"] <= date_to)] if not slips.empty else slips
    col = w[w["row_type"] == "COL"] if not w.empty else w
    all_col = float(col["amount"].sum()) if not col.empty else 0.0
    after = float(col[col["cr_date"] > cutoff]["amount"].sum()) if not col.empty else 0.0
    inj = col[col["cr_date"] <= cutoff] if not col.empty else col
    unmapped = float(inj[inj["store"] == UNRESOLVED]["amount"].sum()) if not inj.empty else 0.0
    adj_net = 0.0
    if not w.empty:
        resolved = w[(w["store"] != UNRESOLVED) & (w["cr_date"] <= cutoff)]
        adj_net = (float(resolved[resolved["row_type"] == "ADJ.CR."]["adj"].fillna(0).sum())
                   - float(resolved[resolved["row_type"] == "ADJ.DB."]["recovered"].fillna(0).sum()))
    j = float(session.query(func.coalesce(func.sum(StoreDayAgg.cash_received), 0)).filter(
        StoreDayAgg.run_id == run_id, StoreDayAgg.business_date >= date_from,
        StoreDayAgg.business_date <= date_to).scalar() or 0)
    expected = all_col - after - unmapped + adj_net
    return Bridge(round(page_cms, 2), round(all_col - page_cms, 2), round(all_col, 2), round(after, 2),
                  round(unmapped, 2), round(adj_net, 2), round(expected, 2), round(j, 2), round(j - expected, 2), cutoff)


# ---------------------------------------------------------------------------
# The lists that keep nothing invisible
# ---------------------------------------------------------------------------

def credited_later(slips: pd.DataFrame, date_from: str, date_to: str) -> pd.DataFrame:
    """COL slips whose credit is not on day or day+1 - in neither CMS column, still outstanding in CLOSING."""
    cols = ["store", "pickup_code", "slip_no", "dep_date", "cr_date", "lag_days", "amount", "flag"]
    if slips.empty:
        return pd.DataFrame(columns=cols)
    col = slips[(slips["row_type"] == "COL") & (slips["dep_date"] >= date_from) & (slips["dep_date"] <= date_to)].copy()
    col["lag_days"] = (pd.to_datetime(col["cr_date"]) - pd.to_datetime(col["dep_date"])).dt.days
    out = col[~col["lag_days"].isin([0, 1])].copy()
    out["flag"] = out["lag_days"].map(lambda g: "" if pd.isna(g) else (f"T+{int(g)} - exceptional" if g >= LATE_FLAG_DAYS
                                                                        else ("credited BEFORE deposit date" if g < 0 else "")))
    out.loc[out["cr_date"].isna(), "flag"] = "no credit date"
    return out.sort_values(["lag_days", "store", "dep_date"], ascending=[False, True, True])[cols].reset_index(drop=True)


def unresolved_rows(slips: pd.DataFrame, date_from: str, date_to: str) -> pd.DataFrame:
    cols = ["pickup_code", "location", "rows", "amount"]
    if slips.empty:
        return pd.DataFrame(columns=cols)
    w = slips[(slips["row_type"] == "COL") & (slips["store"] == UNRESOLVED)
              & (slips["dep_date"] >= date_from) & (slips["dep_date"] <= date_to)]
    if w.empty:
        return pd.DataFrame(columns=cols)
    return (w.groupby(["pickup_code", "location"], dropna=False).agg(rows=("amount", "size"), amount=("amount", "sum"))
            .reset_index().sort_values("amount", ascending=False)[cols].reset_index(drop=True))


def pseudo_store_totals(slips: pd.DataFrame, date_from: str, date_to: str) -> dict:
    out = {}
    for s in PSEUDO_STORES:
        w = slips[(slips["row_type"] == "COL") & (slips["store"] == s) & (slips["dep_date"] >= date_from)
                  & (slips["dep_date"] <= date_to)] if not slips.empty else slips
        out[s] = (int(len(w)), float(w["amount"].sum()) if len(w) else 0.0)
    return out


def excluded_types(slips: pd.DataFrame, date_from: str, date_to: str) -> pd.DataFrame:
    """ADJ.CR., ADJ.DB. (dated, so per period) and ZPO (no date - cannot be placed in any period)."""
    rows = []
    for t in ("ADJ.CR.", "ADJ.DB."):
        w = slips[(slips["row_type"] == t) & (slips["dep_date"] >= date_from) & (slips["dep_date"] <= date_to)] if not slips.empty else slips
        rows.append({"Type": t, "Rows": int(len(w)), "SLIPAMOUNT": float(w["amount"].sum()) if len(w) else 0.0,
                     "ADJ": float(w["adj"].fillna(0).sum()) if len(w) else 0.0,
                     "RECOVEREDAMOUNT": float(w["recovered"].fillna(0).sum()) if len(w) else 0.0,
                     "Note": "dated - in this period"})
    z = slips[slips["row_type"] == "ZPO"] if not slips.empty else slips
    rows.append({"Type": "ZPO", "Rows": int(len(z)), "SLIPAMOUNT": float(z["amount"].sum()) if len(z) else 0.0,
                 "ADJ": 0.0, "RECOVEREDAMOUNT": 0.0,
                 "Note": "no deposit or credit date - cannot be placed in a period (all months in the files)"})
    return pd.DataFrame(rows)


def store_differences(slips: pd.DataFrame, deposits: pd.DataFrame, date_from: str, date_to: str) -> pd.DataFrame:
    """Stores whose PTC deposits for the period differ from the CMS COL slips of the same deposit dates (any credit
    date). The exception list for the page. UNRESOLVED and BNK are reported beside it, not inside it."""
    d = deposits[(deposits["dep_date"] >= date_from) & (deposits["dep_date"] <= date_to)].groupby("store")["amount"].sum()
    c = (slips[(slips["row_type"] == "COL") & (slips["dep_date"] >= date_from) & (slips["dep_date"] <= date_to)]
         .groupby("store")["amount"].sum()) if not slips.empty else pd.Series(dtype=float)
    df = pd.concat([d.rename("deposit"), c.rename("cms")], axis=1).fillna(0.0)
    df = df[~df.index.isin(PSEUDO_STORES)]
    df["difference"] = (df["deposit"] - df["cms"]).round(2)
    df.index.name = "store"
    return (df[df["difference"].abs() > 0.5].sort_values("difference", key=abs, ascending=False)
            .reset_index())


# ---------------------------------------------------------------------------
# Remarks - on a store over a period, surviving every re-run, with history
# ---------------------------------------------------------------------------

def remarks_for(session, date_from: str, date_to: str) -> dict:
    rows = session.query(CmsRecoRemark).filter(CmsRecoRemark.period_from == date_from,
                                              CmsRecoRemark.period_to == date_to).all()
    return {r.store: (r.remark, r.remark_by, r.remark_at) for r in rows}


def set_remark(session, store: str, date_from: str, date_to: str, remark: str, operator: str, now: datetime = None):
    remark = (remark or "").strip()
    if not remark:
        raise ValueError("A remark cannot be empty.")
    if not operator:
        raise ValueError("A remark must record who made it.")
    now = now or datetime.now()
    row = session.query(CmsRecoRemark).filter_by(store=store, period_from=date_from, period_to=date_to).first()
    if row is None:
        session.add(CmsRecoRemark(store=store, period_from=date_from, period_to=date_to, remark=remark,
                                  remark_by=operator, remark_at=now))
    else:
        row.remark, row.remark_by, row.remark_at = remark, operator, now
    session.add(CmsRecoRemarkEvent(store=store, period_from=date_from, period_to=date_to, remark=remark,
                                   remark_by=operator, remark_at=now))
    session.commit()


def remark_history(session, store: str, date_from: str, date_to: str) -> pd.DataFrame:
    rows = (session.query(CmsRecoRemarkEvent).filter_by(store=store, period_from=date_from, period_to=date_to)
            .order_by(CmsRecoRemarkEvent.remark_at.asc(), CmsRecoRemarkEvent.id.asc()).all())
    return pd.DataFrame([{"When": r.remark_at.strftime("%d/%m/%Y %H:%M"), "By": r.remark_by, "Remark": r.remark}
                         for r in rows], columns=["When", "By", "Remark"])
