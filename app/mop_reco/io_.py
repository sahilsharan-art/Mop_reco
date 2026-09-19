"""Raw file reading — Stage 2.

Every function here is deliberately dumb about business meaning: it reads
bytes into a grid of cells, or locates a header row, or picks a file. What
the columns *mean* is ingest.py's job, one normaliser per source.
"""

from __future__ import annotations

import csv
import logging
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Optional

import openpyxl

from .config import Config

logger = logging.getLogger("mop_reco")
if not logger.handlers:
    # CLAUDE.md section 8: never print() for output — use the logger. This is
    # the one place that configures it; every other module just calls
    # logging.getLogger("mop_reco").
    logs_dir = Path(__file__).resolve().parents[2] / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    handler = logging.FileHandler(logs_dir / "mop_reco.log", encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s"))
    logger.addHandler(handler)
    console = logging.StreamHandler()
    console.setFormatter(logging.Formatter("%(levelname)s: %(message)s"))
    logger.addHandler(console)
    logger.setLevel(logging.INFO)


class IOError_(ValueError):
    """Raised on anything that would otherwise be a silent data problem:
    an unreadable file, a decode failure, a header that can't be found."""


# ---------------------------------------------------------------------------
# read_any — one entry point for every raw format this project sees.
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class ReadResult:
    grid: list          # list[list] — raw cell values, no header interpretation
    file_format: str     # "xlsb" | "xlsx" | "csv" | "zip"
    encoding_used: Optional[str]  # None for binary spreadsheet formats
    sheet_name: Optional[str]
    zip_member: Optional[str]
    sheet_row_counts: dict = None  # {sheet: data_row_count} when several were concatenated


def read_any(path: Path, encoding: Optional[str] = None, sheet_name: Optional[str] = None,
             max_rows: Optional[int] = None) -> ReadResult:
    """Read `path` into a raw grid. `encoding` must be supplied by the caller
    for text formats (csv / zip-of-csv) — this function never guesses one.
    Fails loudly on a decode error rather than substituting characters.

    `max_rows` stops the read early (header row included). Use it for a peek —
    a header/shape check must never stream a whole 222MB workbook."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".xlsb":
        return _read_xlsb(path, sheet_name, max_rows)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet_name, max_rows)
    if suffix == ".csv":
        if encoding is None:
            raise IOError_(f"{path.name}: no encoding declared — encoding must come from config, never guessed")
        return _read_csv(path, encoding, max_rows)
    if suffix == ".zip":
        if encoding is None:
            raise IOError_(f"{path.name}: no encoding declared for the CSV inside this zip")
        return _read_zip_csv(path, encoding, max_rows)

    raise IOError_(f"{path.name}: unsupported file type '{suffix}'")


def _read_xlsb(path: Path, sheet_name: Optional[str], max_rows: Optional[int] = None) -> ReadResult:
    """With no sheet_name, reads EVERY sheet and concatenates.

    A monthly POS export exceeds Excel's 1,048,576-row sheet limit and is split
    across `Sheet1`, `Sheet1 (2)`, … — taking sheets[0] silently drops the rest
    (47% of July). Headers must be identical across sheets or this fails loudly
    rather than concatenating mismatched columns. Per-sheet row counts are
    returned so a lost sheet shows up as a number. CLAUDE.md section 5.
    """
    import pyxlsb

    with pyxlsb.open_workbook(str(path)) as wb:
        if sheet_name is not None:
            if sheet_name not in wb.sheets:
                raise IOError_(f"{path.name}: sheet '{sheet_name}' not found (sheets present: {wb.sheets})")
            sheets_to_read = [sheet_name]
        else:
            sheets_to_read = list(wb.sheets)

        grid = []
        header = None
        per_sheet = {}
        for name in sheets_to_read:
            with wb.get_sheet(name) as sheet:
                count = 0
                for row in sheet.rows():
                    vals = [cell.v for cell in row]
                    if count == 0:
                        this_header = [str(v).strip() if v is not None else "" for v in vals]
                        if header is None:
                            header = this_header
                            grid.append(vals)
                        elif this_header != header:
                            raise IOError_(
                                f"{path.name}: sheet '{name}' header differs from the first sheet's — "
                                f"refusing to concatenate mismatched columns.\n"
                                f"  first sheet: {header}\n  '{name}': {this_header}"
                            )
                        count += 1
                        continue
                    grid.append(vals)
                    count += 1
                    if max_rows is not None and len(grid) >= max_rows:
                        break
                per_sheet[name] = count - 1  # exclude that sheet's header row
            if max_rows is not None and len(grid) >= max_rows:
                break

    total = sum(per_sheet.values())
    logger.info(
        f"read_any: {path.name} [xlsb sheets={list(per_sheet)}] — {total} data rows "
        f"(per sheet: {per_sheet})"
    )
    return ReadResult(
        grid=grid, file_format="xlsb", encoding_used=None,
        sheet_name=",".join(per_sheet), zip_member=None, sheet_row_counts=per_sheet,
    )


def _read_xlsx(path: Path, sheet_name: Optional[str], max_rows: Optional[int] = None) -> ReadResult:
    # read_only=True keeps the underlying zip file handle open for lazy
    # row iteration — never closed before, which is harmless for a raw
    # dump (read-only, never moved) but leaves the file locked on Windows
    # against anything that tries to move/delete it right after (found via
    # the Upload page's store_file(), the first caller to do exactly that).
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        chosen_sheet = sheet_name or wb.sheetnames[0]
        if chosen_sheet not in wb.sheetnames:
            raise IOError_(f"{path.name}: sheet '{chosen_sheet}' not found (sheets present: {wb.sheetnames})")
        ws = wb[chosen_sheet]
        grid = []
        for row in ws.iter_rows(values_only=True):
            grid.append(list(row))
            if max_rows is not None and len(grid) >= max_rows:
                break
    finally:
        wb.close()
    logger.info(f"read_any: {path.name} [xlsx sheet={chosen_sheet}] — {len(grid)} rows")
    return ReadResult(grid=grid, file_format="xlsx", encoding_used=None, sheet_name=chosen_sheet, zip_member=None)


def _decode_lines(raw_bytes: bytes, encoding: str, file_label: str):
    try:
        text = raw_bytes.decode(encoding)
    except UnicodeDecodeError as exc:
        raise IOError_(
            f"{file_label}: failed to decode as '{encoding}' — {exc}. "
            "Refusing to silently substitute characters; check the declared encoding in config.yaml."
        ) from exc
    return text.splitlines()


def _read_csv(path: Path, encoding: str, max_rows: Optional[int] = None) -> ReadResult:
    raw_bytes = path.read_bytes()
    lines = _decode_lines(raw_bytes, encoding, path.name)
    if max_rows is not None:
        lines = lines[:max_rows]
    grid = list(csv.reader(lines))
    logger.info(f"read_any: {path.name} [csv encoding={encoding}] — {len(grid)} rows")
    return ReadResult(grid=grid, file_format="csv", encoding_used=encoding, sheet_name=None, zip_member=None)


def _read_zip_csv(path: Path, encoding: str, max_rows: Optional[int] = None) -> ReadResult:
    with zipfile.ZipFile(path) as zf:
        csv_members = [n for n in zf.namelist() if n.lower().endswith(".csv")]
        if len(csv_members) != 1:
            raise IOError_(
                f"{path.name}: expected exactly one .csv member, found {len(csv_members)} "
                f"({csv_members}) — never guessing which one to use"
            )
        member = csv_members[0]
        raw_bytes = zf.read(member)
    lines = _decode_lines(raw_bytes, encoding, f"{path.name}::{member}")
    if max_rows is not None:
        lines = lines[:max_rows]
    grid = list(csv.reader(lines))
    logger.info(f"read_any: {path.name}::{member} [csv-in-zip encoding={encoding}] — {len(grid)} rows")
    return ReadResult(grid=grid, file_format="zip", encoding_used=encoding, sheet_name=None, zip_member=member)


# ---------------------------------------------------------------------------
# find_header_row — locate the header rather than assuming row 1
# ---------------------------------------------------------------------------

def find_header_row(grid: list, required_fields: list, scan_limit: int = 25) -> int:
    """Returns the 0-based index of the first row containing every name in
    `required_fields` as an exact cell value. Scans at most `scan_limit` rows
    (title/blank rows, like the PTC CSV's, sit above the real header)."""
    required = set(required_fields)
    for idx, row in enumerate(grid[:scan_limit]):
        cells = {str(c).strip() for c in row if c is not None}
        if required.issubset(cells):
            return idx
    raise IOError_(
        f"Could not find a header row containing all of {sorted(required)} "
        f"in the first {min(scan_limit, len(grid))} rows"
    )


# ---------------------------------------------------------------------------
# latest_file — most recently modified real data file for a source
# ---------------------------------------------------------------------------

def latest_file(config: Config, source: str, business_date=None):
    """Most recently modified real data file for `source`, excluding lock/temp/
    dotfiles (config.py already excludes them from source_status().files —
    this re-derives via the same _skip_reason so the exclusion rule lives in
    one place). `business_date` is accepted for forward compatibility with the
    per-date subfolder layout in CLAUDE.md section 4; today's folders have no
    such subfolders, so it is not yet used to filter — logged, not silently
    ignored, so this becomes visible the day it matters."""
    status = config.source_status(source)
    if not status.files:
        raise IOError_(f"No usable files found for source '{source}' (path: {status.configured_path})")

    chosen = max(status.files, key=lambda p: p.stat().st_mtime)
    mtime = datetime.fromtimestamp(chosen.stat().st_mtime)
    note = ""
    if business_date is not None:
        note = (
            f" (business_date={business_date} requested, but no per-date subfolder "
            "exists yet — picked most-recently-modified across the whole source folder)"
        )
    logger.info(f"latest_file: source='{source}' chose '{chosen.name}' mtime={mtime}{note}")
    return chosen, mtime


# ---------------------------------------------------------------------------
# clean_ref — CLAUDE.md section 6.2
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Date parsing — the same logical field arrives in four different shapes
# ---------------------------------------------------------------------------

EXCEL_EPOCH = datetime(1899, 12, 30)

# Day-first, always. Never rely on a parser's guess: an unqualified parse of
# '01-07-2026' yields 7 January, and silently NaTs '13-07-2026' and
# '31-07-2026' outright — verified. CLAUDE.md section 6.10.
_DATE_FORMATS = ("%d-%m-%Y", "%d-%b-%Y", "%d-%m-%y", "%d-%b-%y", "%Y-%m-%d")


def parse_business_date(value):
    """One date value -> datetime.date, or None if genuinely blank.

    Handles every shape seen in the real files:
      - Excel serial float from .xlsb   (46223.0  -> 2026-07-20)
      - real datetime from .xlsx        (2026-07-01 00:00:00)
      - day-first string from CSV       ('01-07-2026', '20-Jul-2026')
      - string with a time appended     ('01-07-2026 22:32:38')
    Raises on a non-blank value it cannot parse — never returns a wrong date.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if value != value:  # NaN
            return None
        # Excel serials for this project are ~46,000 (2026). A plain year would
        # be ~2026, so anything that small is not a serial and is a data error.
        if value < 20000:
            raise IOError_(f"suspicious numeric date {value!r} — too small to be an Excel serial")
        return (EXCEL_EPOCH + timedelta(days=float(value))).date()

    s = str(value).strip()
    if s == "" or s.lower() in ("none", "nan", "nat"):
        return None
    s = s.split(" ")[0]  # drop any time portion
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise IOError_(f"could not parse date {value!r} with any of {_DATE_FORMATS}")


def parse_business_date_series(values):
    """parse_business_date across an iterable, for a DataFrame column."""
    return [parse_business_date(v) for v in values]


def parse_bill_time(value):
    """One BILL_TIME value -> 'HH:MM:SS' string, or None if genuinely blank.

    BILL_TIME arrives from the POS `.xlsb` export as an Excel TIME serial —
    a float in [0, 1) that is a fraction of a 24-hour day (0.5 = noon),
    the same serial-number convention BILLDATE uses for the date part, just
    fractional instead of whole. Also handles a real datetime.time or
    datetime.datetime, in case a future export carries it differently, and
    a plain 'HH:MM:SS' string. Raises on a non-blank value it cannot parse
    — never returns a wrong time.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        if value != value:  # NaN
            return None
        if not (0 <= value < 1):
            raise IOError_(f"suspicious BILL_TIME {value!r} — not a fraction of a day")
        total_seconds = round(value * 86400)
        # Rounding a value just under 1.0 (e.g. 0.9999999) can reach 86400
        # seconds — the last instant of the day, not a real 24:00:00. Clamp
        # rather than let it read as midnight the NEXT day.
        total_seconds = min(total_seconds, 86399)
        hh, rem = divmod(total_seconds, 3600)
        mm, ss = divmod(rem, 60)
        return f"{hh:02d}:{mm:02d}:{ss:02d}"
    s = str(value).strip()
    if s == "" or s.lower() in ("none", "nan", "nat"):
        return None
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).strftime("%H:%M:%S")
        except ValueError:
            continue
    raise IOError_(f"could not parse BILL_TIME {value!r} with any known format")


def parse_bill_time_series(values):
    """parse_bill_time across an iterable, for a DataFrame column."""
    return [parse_bill_time(v) for v in values]


_BLANK_REF_VALUES = {"", "0", "nan", "#n/a", "0x2a"}


def clean_ref(value) -> Optional[str]:
    """Strip a leading apostrophe and a trailing '.0'; treat "", 0, nan, #N/A
    and 0x2a (an Excel-error artifact from .xlsb) as blank. Returns None for
    blank, else the cleaned string."""
    if value is None:
        return None
    s = str(value).strip()
    if s.startswith("'"):
        s = s[1:]
    if s.endswith(".0"):
        s = s[:-2]
    if s.strip().lower() in _BLANK_REF_VALUES:
        return None
    return s
